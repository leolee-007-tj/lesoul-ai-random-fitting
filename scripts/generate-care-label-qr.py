#!/usr/bin/env python3
"""Generate deterministic Care Labels QR deliverables for the final Manus domain."""
from __future__ import annotations

import argparse
import csv
import json
import shutil
import zipfile
from pathlib import Path

import qrcode
from qrcode.constants import ERROR_CORRECT_M
from qrcode.image.svg import SvgPathImage
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

FINAL_ORIGIN = "https://lesoulai-xabhv8qd.manus.space"

def esc_typst(value: str) -> str:
    return value.replace("\\", "\\\\").replace('"', '\\"').replace("#", "\\#").replace("$", "\\$")

def product_url(product_id: str) -> str:
    return f"{FINAL_ORIGIN}/product/{product_id}"

def build_qr(url: str):
    qr = qrcode.QRCode(version=None, error_correction=ERROR_CORRECT_M, box_size=16, border=4)
    qr.add_data(url)
    qr.make(fit=True)
    return qr

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    products = json.loads(args.data.read_text(encoding="utf-8"))
    if len(products) != 384:
        raise SystemExit(f"Expected 384 products, got {len(products)}")
    ids = [item["id"] for item in products]
    if ids != [f"LS-{index:04d}" for index in range(1, 385)]:
        raise SystemExit("Product IDs are incomplete or out of order")

    out = args.output.resolve()
    png_dir, svg_dir = out / "PNG", out / "SVG"
    shutil.rmtree(out, ignore_errors=True)
    png_dir.mkdir(parents=True)
    svg_dir.mkdir(parents=True)
    manifest = []
    for item in products:
        pid, name = item["id"], str(item.get("name") or "")
        url = product_url(pid)
        qr = build_qr(url)
        qr.make_image(fill_color="black", back_color="white").save(png_dir / f"{pid}.png")
        qr.make_image(image_factory=SvgPathImage).save(svg_dir / f"{pid}.svg")
        manifest.append({"商品编号": pid, "商品名称": name, "商品页面": url, "PNG文件": f"PNG/{pid}.png", "SVG文件": f"SVG/{pid}.svg"})

    (out / "QR_URL_MANIFEST.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    with (out / "QR_URL_MANIFEST.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(manifest[0]))
        writer.writeheader(); writer.writerows(manifest)

    book = Workbook(); sheet = book.active; sheet.title = "商品二维码匹配表"
    headers = list(manifest[0]); sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="28211C"); cell.alignment = Alignment(horizontal="center")
    for row in manifest: sheet.append([row[key] for key in headers])
    for width, column in zip([14, 30, 62, 24, 24], "ABCDE"):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = f"A1:E{len(manifest)+1}"
    book.save(out / "QR_MATCHING.xlsx")

    typ = [
        '#set page(paper: "a4", margin: 8mm)',
        '#set text(font: ("Noto Sans CJK SC", "Noto Sans CJK KR"), lang: "zh", size: 7pt, fill: rgb("28211c"))',
        '#set par(leading: 0.75em)',
        '#let card(id, name, url, qr) = box(width: 46mm, height: 52mm, stroke: .35pt + rgb("cfc7bd"), radius: 2mm, inset: 2.5mm)[',
        '  #align(center)[',
        '    #text(size: 6.5pt, weight: "bold", tracking: .8pt, fill: rgb("74685d"))[LESOUL 商品二维码]',
        '    #v(1.2mm)',
        '    #box(width: 41mm, height: 7.5mm, clip: true)[#align(center + horizon)[#text(size: 8.5pt, weight: "bold")[#name]]]',
        '    #text(size: 6.5pt, weight: "bold", fill: rgb("887b70"))[#id]',
        '    #v(1mm)',
        '    #image(qr, width: 27mm, height: 27mm)',
        '    #v(.8mm)',
        '    #text(size: 5.8pt)[微信或手机相机扫码查看商品洗标资料]',
        '    #v(.5mm)',
        '    #box(width: 40mm, height: 3mm, clip: true)[#text(size: 3.7pt, fill: rgb("9a8f85"))[#url]]',
        '  ]',
        ']',
    ]
    for page_start in range(0, len(products), 20):
        page = products[page_start:page_start+20]
        typ.append('#grid(columns: (46mm, 46mm, 46mm, 46mm), column-gutter: 3mm, row-gutter: 3mm,')
        for item in page:
            pid, name = item["id"], str(item.get("name") or "")
            url = product_url(pid)
            typ.append(f'  card("{pid}", "{esc_typst(name)}", "{url}", "SVG/{pid}.svg"),')
        typ.append(')')
        if page_start + 20 < len(products): typ.append('#pagebreak()')
    (out / "LESOUL_QR_CARDS_CHINESE.typ").write_text("\n".join(typ) + "\n", encoding="utf-8")

    zip_path = out / "LESOUL_QR_PNG_SVG_FINAL.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(png_dir.glob("*.png")) + sorted(svg_dir.glob("*.svg")):
            archive.write(path, path.relative_to(out))
        for name in ["QR_URL_MANIFEST.json", "QR_URL_MANIFEST.csv", "QR_MATCHING.xlsx"]:
            archive.write(out / name, name)
    print(json.dumps({"products": len(products), "origin": FINAL_ORIGIN, "output": str(out)}, ensure_ascii=False))

if __name__ == "__main__":
    main()
